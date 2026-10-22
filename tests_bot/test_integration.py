from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select

import avito_client
import link_resolver
from intelligence.analysis import (
    Market,
    assess,
    condition_key,
    defects,
    model_key,
    seller_type,
)
from intelligence.config import Settings
from intelligence.db import Listing, PriceHistory, Search, session_factory
from intelligence.export import export_search
from intelligence.service import Service, validate_url
from intelligence.vision import fingerprint

URL = "https://www.avito.ru/moskva/telefony?q=iphone+15"


def item(
    ident,
    title="iPhone 15, 128 ГБ",
    price=40000,
    description="Хорошее состояние",
    seller="Частное лицо",
    closed=2,
):
    return {
        "id": ident,
        "sellerId": ident + 1000,
        "title": title,
        "priceDetailed": {"value": price},
        "description": description,
        "urlPath": f"/moskva/telefony/iphone_{ident}",
        "closedItemsText": f"{closed} завершённых объявлений",
        "iva": {"SecondLineStep": [{"payload": {"value": seller}}]},
    }


class FakeSource:
    def __init__(self):
        self.pages = {}
        self.calls = []
        self.fail = False

    def resolve(self, url):
        return link_resolver.ResolvedSearch(
            "iPhone 15 Москва",
            "goods",
            url,
            {"q": "iphone 15"},
            "Москва",
            30,
            ["Запрос: iPhone 15"],
            [],
            [],
        )

    def page(self, search, number):
        if self.fail:
            raise RuntimeError("источник недоступен")
        self.calls.append((search.id, number))
        rows = self.pages.get((search.id, number), [])
        if isinstance(rows, avito_client.Page):
            return rows
        return avito_client.Page(number, tuple(rows), len(rows), len(rows))


@pytest.fixture
def system(tmp_path):
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'bot.sqlite'}",
        telegram_allowed_user_ids="1,2",
        check_interval_seconds=300,
        request_delay_seconds=10,
    )
    sessions = session_factory(settings.database_url)
    source = FakeSource()
    return Service(sessions, source, settings), source, sessions, tmp_path


def test_full_lifecycle(system):
    service, source, sessions, tmp = system
    search = service.create_search(
        1, URL, target_model="iPhone 15", target_memory="128 ГБ"
    )
    service.set_filter(1, search.id, "min_score", "60")
    source.pages[(search.id, 1)] = [item(n) for n in range(1, 23)] + [
        item(23, price=1000000),
        item(24, title="iPhone 15 Pro, 128 ГБ", price=45000),
        item(25, description="Не работает, на запчасти", price=5000),
        item(26, seller="Магазин", closed=500),
        item(27, price=39000),
        item(28, price=41000),
        item(29, price=42000),
        item(30, price=43000),
    ]
    result = service.scan_one(search.id)
    assert result["complete"] and result["new"] == 30
    assert not service.pending()  # baseline is silent
    market = service.stats(1, search.id)["market"]
    assert market.median == 40000
    assert market.count >= 20
    source.pages[(search.id, 1)] = [item(31, price=32000)] + source.pages[
        (search.id, 1)
    ]
    service.scan_one(search.id)
    assert len(service.pending()) == 1
    first = service.pending()[0][0]
    assert first.listing_id == "31"
    service.mark_sent(first.id, True)
    service.scan_one(search.id)
    assert not service.pending()
    source.pages[(search.id, 1)][0] = item(31, price=30000)
    service.scan_one(search.id)
    assert len(service.pending()) == 1
    service.favorite(1, "31", False)
    assert not service.pending()
    assert service.listing_details(1, "31")[0].exclusion == "неинтересно пользователю"
    service.favorite(1, "31", True)
    assert service.favorites(1)[0].id == "31"
    service.exclude_seller(1, "31")
    assert service.listing_details(1, "31")[0].exclusion == "продавец исключён"
    with sessions() as db:
        assert (
            len(
                list(
                    db.scalars(
                        select(PriceHistory).where(PriceHistory.listing_id == "31")
                    )
                )
            )
            == 2
        )
    service.set_status(1, search.id, "paused")
    assert search.id not in service.scan_due()
    service2 = Service(sessions, source, service.settings)
    assert service2.list_searches(1)[0].status == "paused"
    service2.set_status(1, search.id, "active")
    assert service2.list_searches(1)[0].status == "active"
    second = service2.create_search(
        2, "https://www.avito.ru/kazan/telefony?q=steam+deck"
    )
    assert service2.stats(2, second.id)["total"] == 0
    with pytest.raises(ValueError):
        service2.stats(2, search.id)
    path = export_search(sessions, 1, search.id, tmp / "export.xlsx")
    book = load_workbook(path)
    assert set(book.sheetnames) == {
        "Все объявления",
        "Выгодные",
        "Исключённые",
        "История цен",
        "Анализ рынка",
    }
    assert book["Все объявления"].max_row == 32
    assert book["История цен"].max_row >= 33


def test_source_error_and_resume(system):
    service, source, _, _ = system
    search = service.create_search(1, URL)
    source.fail = True
    assert "error" in service.scan_one(search.id)
    assert service.list_searches(1)[0].status == "source_unavailable"
    assert service.list_searches(1)[0].failure_count == 1
    assert "error" in service.scan_one(search.id)
    assert service.list_searches(1)[0].failure_count == 2
    source.fail = False
    service.scan_one(search.id)
    assert service.list_searches(1)[0].status == "active"
    assert service.list_searches(1)[0].failure_count == 0


@pytest.mark.parametrize("fails", [False, True])
def test_pause_during_source_request_is_preserved(system, fails):
    service, source, sessions, _ = system
    search = service.create_search(1, URL)

    def page(_search, _number):
        service.set_status(1, search.id, "paused")
        if fails:
            raise RuntimeError("источник недоступен")
        return avito_client.Page(1, (item(1),), 1, 1)

    source.page = page
    service.scan_one(search.id)
    with sessions() as db:
        saved = db.get(Search, search.id)
        assert saved.status == "paused"
        assert saved.scan_complete is False
        assert service.stats(1, search.id)["total"] == 0


def test_watch_multiple_pages_and_reappearance(system):
    service, source, sessions, _ = system
    service.settings.watch_pages = 2
    search = service.create_search(1, URL)
    source.pages[(search.id, 1)] = [item(1), item(2)]
    service.scan_one(search.id)
    source.pages[(search.id, 1)] = avito_client.Page(
        1, tuple(item(i) for i in range(3, 53)), 52, 50
    )
    source.pages[(search.id, 2)] = avito_client.Page(2, (item(53),), 52, 1)
    service.scan_one(search.id)
    with sessions() as db:
        assert db.get(Search, search.id).watch_page == 2
    service.scan_one(search.id)
    with sessions() as db:
        assert db.get(Search, search.id).watch_page == 1
        from intelligence.db import Match

        assert db.get(Match, (search.id, "1")).state == "not_seen"
    source.pages[(search.id, 1)] = [item(1)]
    service.scan_one(search.id)
    with sessions() as db:
        assert db.get(Match, (search.id, "1")).state == "reappeared"


def test_analysis_and_input_validation():
    assert model_key("iPhone 15, 128 ГБ") != model_key("iPhone 15 Pro, 128 ГБ")
    assert condition_key("iPhone 15 новый") != condition_key("iPhone 15 б/у")
    assert defects("Никаких проблем с Face ID нет") == []
    assert defects("Face ID не работает")
    assert seller_type("Магазин", "Наш магазин, работаем ежедневно", 150).kind == "shop"
    assert seller_type("", "", None).kind == "unknown"
    assert (
        seller_type("Частное лицо", "Я не магазин, продаю свой телефон", 2).kind
        == "likely_private"
    )
    for url in [
        "http://www.avito.ru/moskva",
        "https://evilavito.ru/moskva",
        "https://www.avito.ru.evil.com/moskva",
    ]:
        with pytest.raises(ValueError):
            validate_url(url)


def test_duplicate_photo_warning(system):
    service, source, _, _ = system
    search = service.create_search(1, URL)
    source.pages[(search.id, 1)] = [item(1), item(2)]
    service.scan_one(search.id)
    image = Image.new("RGB", (60, 60), "red")
    output = io.BytesIO()
    image.save(output, format="PNG")
    hash_value = fingerprint(output.getvalue()).perceptual_hash
    service.save_image_analysis("1", [hash_value], [], "Обработано локально", None)
    service.save_image_analysis("2", [hash_value], [], "Обработано локально", None)
    assert "похожа" in service.cached_image_analysis("2").warnings[0]
    score = service.listing_details(1, "2")[0].score
    service.save_image_analysis("2", [hash_value], [], "Обработано локально", None)
    assert service.listing_details(1, "2")[0].score == score


def test_discount_math_and_defect_exclusion(system):
    service, source, sessions, _ = system
    search = service.create_search(
        1, URL, target_model="iPhone 15", target_memory="128 ГБ"
    )
    source.pages[(search.id, 1)] = [
        item(1, price=34000),
        item(2, price=5000, description="Не работает"),
    ]
    service.scan_one(search.id)
    with sessions() as db:
        listing = db.get(Listing, "1")
        broken = db.get(Listing, "2")
        saved_search = db.get(Search, search.id)
        market = Market("iphone 15|128gb|unknown", 40000, 38000, 42000, 10, "medium")
        assert assess(listing, saved_search, market).discount_percent == 15
        assert assess(broken, saved_search, market).exclusion


def test_condition_filter(system):
    service, source, sessions, _ = system
    search = service.create_search(1, URL)
    source.pages[(search.id, 1)] = [
        item(1, description="Новый, запечатан"),
        item(2, description="Б/у, пользовались"),
    ]
    service.scan_one(search.id)
    service.set_filter(1, search.id, "condition_filter", "new")
    with sessions() as db:
        from intelligence.db import Match

        assert not db.get(Match, (search.id, "1")).exclusion
        assert db.get(Match, (search.id, "2")).exclusion


def test_quiet_hours_persist(system):
    service, source, sessions, _ = system
    service.set_quiet_hour(1, "start", "22")
    service.set_quiet_hour(1, "end", "7")
    assert Service(sessions, source, service.settings).user_settings(1) == (22, 7)
    service.set_quiet_hour(1, "start", "-")
    assert service.user_settings(1) == (None, 7)
    with pytest.raises(ValueError):
        service.set_quiet_hour(1, "end", "25")


def test_digest_preferences_and_queue(system):
    from datetime import timedelta

    from intelligence.db import User, now

    service, source, sessions, _ = system
    search = service.create_search(
        1, URL, target_model="iPhone 15", target_memory="128 ГБ"
    )
    service.set_filter(1, search.id, "min_score", "50")
    source.pages[(search.id, 1)] = [item(n) for n in range(1, 10)]
    service.scan_one(search.id)
    source.pages[(search.id, 1)] = [item(10, price=33000)] + source.pages[
        (search.id, 1)
    ]
    service.scan_one(search.id)
    notes = service.pending()
    assert notes
    assert service.toggle_digest(1)
    assert not service.digest_due(1)
    with sessions.begin() as db:
        db.get(User, 1).last_digest_at = now() - timedelta(hours=13)
    assert service.digest_due(1)
    service.mark_digest_sent(1, [notes[0][0].id])
    assert not service.pending()
    assert not service.toggle_digest(1)


def test_filter_change_cancels_pending(system):
    service, source, _, _ = system
    search = service.create_search(
        1, URL, target_model="iPhone 15", target_memory="128 ГБ"
    )
    service.set_filter(1, search.id, "min_score", "50")
    source.pages[(search.id, 1)] = [item(n) for n in range(1, 11)]
    service.scan_one(search.id)
    source.pages[(search.id, 1)] = [item(11, price=33000)] + source.pages[
        (search.id, 1)
    ]
    service.scan_one(search.id)
    assert service.pending()
    service.set_filter(1, search.id, "min_score", "100")
    assert not service.pending()
    service.set_filter(1, search.id, "max_price", "50000")
    service.set_filter(1, search.id, "max_price", "-")
    assert service.list_searches(1)[0].max_price is None


def test_bounded_client_stops_on_source_block(monkeypatch):
    import intelligence.service as source_module
    from intelligence.service import BoundedClient

    client = BoundedClient(delay=10, log=lambda *args, **kwargs: None)
    monkeypatch.setattr(
        source_module.main.requests, "Session", lambda **kwargs: object()
    )
    monkeypatch.setattr(source_module.main, "set_session_headers", lambda *args: None)
    monkeypatch.setattr(client, "_get", lambda *args, **kwargs: ("blocked", None))
    with pytest.raises(RuntimeError, match="ограничил доступ"):
        client._open_session(URL)
    client.stopping.set()
    with pytest.raises(RuntimeError, match="остановлен"):
        client._pause()
