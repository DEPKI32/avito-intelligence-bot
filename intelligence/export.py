from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .db import Listing, MarketSnapshot, Match, PriceHistory, Search


def safe(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def export_search(sessions, user_id: int, search_id: int, output: Path) -> Path:
    with sessions() as db:
        search = db.get(Search, search_id)
        if not search or search.user_id != user_id:
            raise ValueError("Поиск не найден")
        rows = list(
            db.execute(
                select(Match, Listing)
                .join(Listing, Match.listing_id == Listing.id)
                .where(Match.search_id == search_id)
            )
        )
        prices = (
            list(
                db.scalars(
                    select(PriceHistory).where(
                        PriceHistory.listing_id.in_([listing.id for _, listing in rows])
                    )
                )
            )
            if rows
            else []
        )
        markets = list(
            db.scalars(
                select(MarketSnapshot)
                .where(MarketSnapshot.search_id == search_id)
                .order_by(MarketSnapshot.calculated_at.desc())
            )
        )
    book = Workbook()
    book.remove(book.active)
    columns = [
        "ID",
        "Название",
        "Цена, ₽",
        "Рынок, ₽",
        "Разница, ₽",
        "Разница, %",
        "Оценка",
        "Риск",
        "Исключение",
        "Причины",
        "Продавец",
        "Первое обнаружение",
        "Последняя проверка",
        "Ссылка",
    ]

    def sheet(name, header, data):
        page = book.create_sheet(name)
        page.append(header)
        for entry in data:
            page.append([safe(x) for x in entry])
        page.freeze_panes = "A2"
        page.auto_filter.ref = page.dimensions
        for cell in page[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24527A")
        for idx, heading in enumerate(header, 1):
            page.column_dimensions[get_column_letter(idx)].width = min(
                55, max(14, len(heading) + 3)
            )
        return page

    def row(match, listing):
        difference = (
            match.market_price - listing.price
            if match.market_price and listing.price
            else None
        )
        percentage = difference / match.market_price if difference is not None else None
        return [
            listing.id,
            listing.title,
            listing.price,
            match.market_price,
            difference,
            percentage,
            match.score,
            match.risk,
            match.exclusion,
            "; ".join(match.reasons or []),
            listing.seller_label,
            match.first_seen.replace(tzinfo=None),
            match.last_seen.replace(tzinfo=None),
            listing.url,
        ]

    sheet("Все объявления", columns, [row(m, l) for m, l in rows])
    sheet(
        "Выгодные",
        columns,
        [
            row(m, l)
            for m, l in rows
            if not m.exclusion and m.market_price and (m.score or 0) >= search.min_score
        ],
    )
    sheet("Исключённые", columns, [row(m, l) for m, l in rows if m.exclusion])
    sheet(
        "История цен",
        ["ID", "Цена, ₽", "Наблюдалась"],
        [[p.listing_id, p.price, p.observed_at.replace(tzinfo=None)] for p in prices],
    )
    sheet(
        "Анализ рынка",
        [
            "Группа",
            "Медиана, ₽",
            "Типичный минимум, ₽",
            "Типичный максимум, ₽",
            "Выборка",
            "Доверие",
            "Рассчитано",
        ],
        [
            [
                m.group_key,
                m.median,
                m.low,
                m.high,
                m.count,
                m.confidence,
                m.calculated_at.replace(tzinfo=None),
            ]
            for m in markets
        ],
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    return output
